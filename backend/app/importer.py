"""Reads QIST's member spreadsheets (the 'R&D Expert List' workbook and similar) into claimable profile records.

Rules, in order of importance:
- Only people with a usable e-mail address become records: the e-mail is how an invitation reaches them and how a claim is
  matched. Rows without one are counted and reported so a person can deal with them by hand; they are never guessed at.
- Only the columns a profile needs are read. Notes, statements, hobbies and any other column are ignored on purpose.
- The same e-mail on several rows (a person in two sheets) becomes one record; the fullest values win.
- Nothing here writes to the database; the caller decides what to do with the records.
"""
import io
import re
from typing import Iterable

from openpyxl import load_workbook

# header (lower-case, trimmed) -> what it means. Several spellings occur across the sheets.
HEADERS = {
    "name surname": "name", "full name": "name", "name": "first", "surname": "last", "last name": "last", "first name": "first",
    "email": "email", "e-mail": "email", "e mail": "email",
    "primary field": "field", "field": "field", "primary field 2": "field2",
    "keywords": "keywords", "research topic": "keywords", "reasech topic": "keywords", "research topics": "keywords",
    "institution": "institution", "organization": "institution", "organisation": "institution",
    "position": "title", "academic role": "title",
    "country work": "country", "country": "country", "city": "city",
    "linkedin": "linkedin", "social media": "social",
}
EMAIL = re.compile(r"[A-Za-z0-9._%+'-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}")


def _text(v) -> str:
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def _first_email(*cells) -> tuple[str, str]:
    """-> (email, why_not). A cell can hold several addresses or surrounding text; the first valid one is taken."""
    seen_any = False
    for c in cells:
        t = _text(c)
        if not t:
            continue
        seen_any = True
        m = EMAIL.search(t)
        if m:
            return m.group(0).lower(), ""
    return "", ("bad_email" if seen_any else "no_email")


def _split_keywords(s: str) -> list[str]:
    out = []
    for k in re.split(r"[;,\n|/]+", s):
        k = k.strip(" .-•")
        if 2 <= len(k) <= 60 and k.lower() not in {x.lower() for x in out}:
            out.append(k)
    return out[:12]


def _linkedin(*cells) -> str:
    for c in cells:
        m = re.search(r"https?://[^\s,;]*linkedin\.com[^\s,;]*", _text(c), re.I)
        if m:
            return m.group(0).rstrip(".)")
    return ""


def _rows(ws) -> Iterable[dict]:
    it = ws.iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        return
    cols = {}
    for i, h in enumerate(header):
        key = HEADERS.get(_text(h).lower())
        if key:
            cols.setdefault(key, []).append(i)
    if "email" not in cols and "name" not in cols and "first" not in cols:
        return                                        # not a member sheet (e.g. an empty tab)
    for row in it:
        if not row or not any(c not in (None, "") for c in row):
            continue
        get = lambda k: [row[i] for i in cols.get(k, []) if i < len(row)]  # noqa: E731
        name = _text((get("name") or [""])[0])
        if not name:
            name = " ".join(x for x in (_text((get("first") or [""])[0]), _text((get("last") or [""])[0])) if x)
        yield {"name": name, "email_cells": get("email"), "field": " ".join(_text(x) for x in get("field") + get("field2")),
               "keywords": "; ".join(_text(x) for x in get("keywords") if _text(x)), "institution": _text((get("institution") or [""])[0]),
               "title": _text((get("title") or [""])[0]), "country": _text((get("country") or [""])[0]), "city": _text((get("city") or [""])[0]),
               "linkedin": _linkedin(*get("linkedin"), *get("social"))}


def parse_workbook(data: bytes, classify, country_code, suppressed: set[str] = frozenset(), existing: set[str] = frozenset()) -> tuple[list[dict], dict]:
    """-> (records, report). `classify(text, body) -> [field]`, `country_code(text) -> ISO or ''` are passed in so this module
    stays free of the rest of the app. `suppressed` and `existing` are sets of lower-case e-mails to leave alone."""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    report = {"sheets": {}, "rows": 0, "no_email": [], "bad_email": [], "no_name": 0, "merged": 0, "already_accounts": 0, "suppressed": 0}
    by_email: dict[str, dict] = {}
    for ws in wb.worksheets:
        n = 0
        for r in _rows(ws):
            n += 1
            report["rows"] += 1
            email, why = _first_email(*r["email_cells"])
            if not r["name"] or len(r["name"]) < 2:
                report["no_name"] += 1
                continue
            if why:
                (report["no_email"] if why == "no_email" else report["bad_email"]).append(r["name"])
                continue
            if email in suppressed:
                report["suppressed"] += 1
                continue
            if email in existing:
                report["already_accounts"] += 1
                continue
            rec = by_email.get(email)
            if rec:
                report["merged"] += 1
                for k in ("name", "title", "institution", "city", "country", "linkedin", "field"):
                    if len(r[k]) > len(rec[k]):
                        rec[k] = r[k]
                rec["keywords"] = "; ".join(x for x in (rec["keywords"], r["keywords"]) if x)
            else:
                by_email[email] = dict(r, email=email)
        report["sheets"][ws.title] = n
    records = []
    for email, r in by_email.items():
        kws = _split_keywords(r["keywords"])
        records.append({"email": email, "name": r["name"][:120], "title": r["title"][:120], "institution": r["institution"][:160],
                        "country": country_code(r["country"]) if r["country"] else "", "city": r["city"][:80], "keywords": kws,
                        "fields": classify(r["field"], " ".join(kws + [r["title"]]))[:6], "linkedin": r["linkedin"]})
    report["records"] = len(records)
    report["no_email"] = sorted(set(report["no_email"]))
    report["bad_email"] = sorted(set(report["bad_email"]))
    return records, report
